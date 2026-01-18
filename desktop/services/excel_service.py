"""
Excel Service - Import dan Export data ke/dari Excel
"""
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from datetime import datetime, date, time
from typing import Optional, List, Dict, Any
import os

from database.models import (
    Jamaah, FaseKehidupan, Wilayah, Jenjang, TahunAjaran,
    Enrollment, AnggotaKelas, Pengajian, JadwalRutin, KelasPengajian,
    KeaktifanPengajian, BidangMateri, KategoriMateri, MateriItem,
    ProgressJamaah, PenilaianAkhlaq, LimaUnsur
)


class ExcelService:
    """Service untuk import dan export data Excel"""

    # Style constants
    HEADER_FILL = PatternFill(start_color="1a5f2b", end_color="1a5f2b", fill_type="solid")
    HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
    HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
    CELL_BORDER = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    def __init__(self, session: Session):
        self.session = session

    def _style_header(self, ws, row: int, num_cols: int):
        """Apply style to header row"""
        for col in range(1, num_cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = self.HEADER_FILL
            cell.font = self.HEADER_FONT
            cell.alignment = self.HEADER_ALIGNMENT
            cell.border = self.CELL_BORDER

    def _auto_width(self, ws):
        """Auto-adjust column widths"""
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

    def _add_data_row(self, ws, row: int, data: List[Any]):
        """Add data row with styling"""
        for col, value in enumerate(data, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = self.CELL_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    # ==================== EXPORT METHODS ====================

    def export_all(self, filepath: str) -> bool:
        """Export semua data ke satu file Excel dengan multiple sheets"""
        try:
            wb = Workbook()
            wb.remove(wb.active)  # Remove default sheet

            # Export each module
            self._export_wilayah_sheet(wb)
            self._export_jenjang_sheet(wb)
            self._export_jamaah_sheet(wb)
            self._export_enrollment_sheet(wb)
            self._export_kurikulum_sheet(wb)
            self._export_pengajian_sheet(wb)
            self._export_presensi_sheet(wb)
            self._export_organisasi_sheet(wb)

            wb.save(filepath)
            return True
        except Exception as e:
            print(f"Error exporting all data: {e}")
            return False

    def export_jamaah(self, filepath: str, wilayah_id: Optional[int] = None) -> bool:
        """Export data jamaah/generus"""
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Data Generus"

            # Header
            headers = [
                "No", "Nama", "Nama Panggilan", "Jenis Kelamin", "Tempat Lahir",
                "Tanggal Lahir", "Umur", "Phone", "Email", "Alamat",
                "Status Pernikahan", "Golongan Darah", "Pendidikan", "Pekerjaan",
                "Status Aktif", "Wilayah", "Jenjang", "Fase Kehidupan"
            ]
            ws.append(headers)
            self._style_header(ws, 1, len(headers))

            # Query data
            query = self.session.query(Jamaah).filter(Jamaah.status_aktif == True)

            if wilayah_id:
                # Filter by wilayah through enrollment
                query = query.join(Enrollment).filter(
                    Enrollment.wilayah_id == wilayah_id,
                    Enrollment.status == 'aktif'
                )

            jamaah_list = query.order_by(Jamaah.nama).all()

            for idx, j in enumerate(jamaah_list, 1):
                # Get active enrollment
                enrollment = self.session.query(Enrollment).filter(
                    Enrollment.jamaah_id == j.id,
                    Enrollment.status == 'aktif'
                ).first()

                wilayah_nama = enrollment.wilayah.nama if enrollment and enrollment.wilayah else ""
                jenjang_nama = enrollment.jenjang.nama if enrollment and enrollment.jenjang else ""

                # Get current fase
                fase = self.session.query(FaseKehidupan).filter(
                    FaseKehidupan.jamaah_id == j.id
                ).order_by(FaseKehidupan.tanggal_masuk.desc()).first()
                fase_nama = fase.fase if fase else ""

                row_data = [
                    idx,
                    j.nama,
                    j.nama_panggilan or "",
                    "Laki-laki" if j.jenis_kelamin == "L" else "Perempuan" if j.jenis_kelamin == "P" else "",
                    j.tempat_lahir or "",
                    j.tanggal_lahir.strftime("%Y-%m-%d") if j.tanggal_lahir else "",
                    j.umur or "",
                    j.phone or "",
                    j.email or "",
                    j.alamat_lengkap or "",
                    j.status_pernikahan or "",
                    j.golongan_darah or "",
                    j.pendidikan_terakhir or "",
                    j.pekerjaan or "",
                    "Aktif" if j.status_aktif else "Tidak Aktif",
                    wilayah_nama,
                    jenjang_nama,
                    fase_nama
                ]
                self._add_data_row(ws, idx + 1, row_data)

            self._auto_width(ws)
            wb.save(filepath)
            return True
        except Exception as e:
            print(f"Error exporting jamaah: {e}")
            return False

    def export_jamaah_per_kelompok(self, filepath: str, kelompok_id: int) -> bool:
        """Export data generus per kelompok"""
        return self.export_jamaah(filepath, wilayah_id=kelompok_id)

    def export_wilayah(self, filepath: str) -> bool:
        """Export data wilayah"""
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Data Wilayah"

            headers = ["No", "Kode", "Nama", "Tingkat", "Parent", "Status"]
            ws.append(headers)
            self._style_header(ws, 1, len(headers))

            wilayah_list = self.session.query(Wilayah).order_by(
                Wilayah.tingkat, Wilayah.nama
            ).all()

            for idx, w in enumerate(wilayah_list, 1):
                parent_nama = w.parent.nama if w.parent else ""
                row_data = [
                    idx,
                    w.kode or "",
                    w.nama,
                    w.tingkat,
                    parent_nama,
                    "Aktif" if w.is_aktif else "Tidak Aktif"
                ]
                self._add_data_row(ws, idx + 1, row_data)

            self._auto_width(ws)
            wb.save(filepath)
            return True
        except Exception as e:
            print(f"Error exporting wilayah: {e}")
            return False

    def export_kurikulum(self, filepath: str) -> bool:
        """Export data kurikulum (bidang, kategori, materi)"""
        try:
            wb = Workbook()

            # Sheet 1: Bidang Materi
            ws1 = wb.active
            ws1.title = "Bidang Materi"
            headers1 = ["No", "Nama Bidang", "Urutan", "Status"]
            ws1.append(headers1)
            self._style_header(ws1, 1, len(headers1))

            bidang_list = self.session.query(BidangMateri).order_by(BidangMateri.urutan).all()
            for idx, b in enumerate(bidang_list, 1):
                self._add_data_row(ws1, idx + 1, [
                    idx, b.nama, b.urutan, "Aktif" if b.is_aktif else "Tidak Aktif"
                ])
            self._auto_width(ws1)

            # Sheet 2: Kategori Materi
            ws2 = wb.create_sheet("Kategori Materi")
            headers2 = ["No", "Bidang", "Kode", "Nama Kategori", "Urutan", "Status"]
            ws2.append(headers2)
            self._style_header(ws2, 1, len(headers2))

            kategori_list = self.session.query(KategoriMateri).order_by(
                KategoriMateri.bidang_id, KategoriMateri.urutan
            ).all()
            for idx, k in enumerate(kategori_list, 1):
                bidang_nama = k.bidang.nama if k.bidang else ""
                self._add_data_row(ws2, idx + 1, [
                    idx, bidang_nama, k.kode or "", k.nama, k.urutan,
                    "Aktif" if k.is_aktif else "Tidak Aktif"
                ])
            self._auto_width(ws2)

            # Sheet 3: Materi Item
            ws3 = wb.create_sheet("Materi Item")
            headers3 = ["No", "Bidang", "Kategori", "Nomor", "Nama Materi", "Tipe", "Status"]
            ws3.append(headers3)
            self._style_header(ws3, 1, len(headers3))

            materi_list = self.session.query(MateriItem).order_by(
                MateriItem.kategori_id, MateriItem.nomor
            ).all()
            for idx, m in enumerate(materi_list, 1):
                bidang_nama = m.kategori.bidang.nama if m.kategori and m.kategori.bidang else ""
                kategori_nama = m.kategori.nama if m.kategori else ""
                self._add_data_row(ws3, idx + 1, [
                    idx, bidang_nama, kategori_nama, m.nomor or "", m.nama, m.tipe,
                    "Aktif" if m.is_aktif else "Tidak Aktif"
                ])
            self._auto_width(ws3)

            wb.save(filepath)
            return True
        except Exception as e:
            print(f"Error exporting kurikulum: {e}")
            return False

    def export_pengajian(self, filepath: str, wilayah_id: Optional[int] = None,
                         tanggal_mulai: Optional[date] = None,
                         tanggal_selesai: Optional[date] = None) -> bool:
        """Export data pengajian"""
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Data Pengajian"

            headers = [
                "No", "Tanggal", "Wilayah", "Jenjang", "Waktu Mulai", "Waktu Selesai",
                "Materi", "Status", "Jumlah Hadir", "Jumlah Total", "Persentase"
            ]
            ws.append(headers)
            self._style_header(ws, 1, len(headers))

            query = self.session.query(Pengajian).order_by(Pengajian.tanggal.desc())

            if wilayah_id:
                query = query.filter(Pengajian.wilayah_id == wilayah_id)
            if tanggal_mulai:
                query = query.filter(Pengajian.tanggal >= tanggal_mulai)
            if tanggal_selesai:
                query = query.filter(Pengajian.tanggal <= tanggal_selesai)

            pengajian_list = query.all()

            for idx, p in enumerate(pengajian_list, 1):
                wilayah_nama = p.wilayah.nama if p.wilayah else ""
                jenjang_nama = p.jenjang.nama if p.jenjang else ""
                row_data = [
                    idx,
                    p.tanggal.strftime("%Y-%m-%d") if p.tanggal else "",
                    wilayah_nama,
                    jenjang_nama,
                    p.waktu_mulai.strftime("%H:%M") if p.waktu_mulai else "",
                    p.waktu_selesai.strftime("%H:%M") if p.waktu_selesai else "",
                    p.materi_pokok or "",
                    p.status,
                    p.jumlah_hadir,
                    p.jumlah_total,
                    f"{p.persentase_hadir:.1f}%"
                ]
                self._add_data_row(ws, idx + 1, row_data)

            self._auto_width(ws)
            wb.save(filepath)
            return True
        except Exception as e:
            print(f"Error exporting pengajian: {e}")
            return False

    def export_presensi(self, filepath: str, pengajian_id: Optional[int] = None,
                        wilayah_id: Optional[int] = None) -> bool:
        """Export data presensi"""
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Data Presensi"

            headers = [
                "No", "Tanggal Pengajian", "Wilayah", "Nama Jamaah",
                "Status", "Keterangan"
            ]
            ws.append(headers)
            self._style_header(ws, 1, len(headers))

            query = self.session.query(KeaktifanPengajian).join(Pengajian).order_by(
                Pengajian.tanggal.desc()
            )

            if pengajian_id:
                query = query.filter(KeaktifanPengajian.pengajian_id == pengajian_id)
            if wilayah_id:
                query = query.filter(Pengajian.wilayah_id == wilayah_id)

            presensi_list = query.all()

            for idx, pr in enumerate(presensi_list, 1):
                tanggal = pr.pengajian.tanggal.strftime("%Y-%m-%d") if pr.pengajian and pr.pengajian.tanggal else ""
                wilayah = pr.pengajian.wilayah.nama if pr.pengajian and pr.pengajian.wilayah else ""
                nama = pr.jamaah.nama if pr.jamaah else ""
                row_data = [
                    idx, tanggal, wilayah, nama,
                    pr.status_display, pr.keterangan or ""
                ]
                self._add_data_row(ws, idx + 1, row_data)

            self._auto_width(ws)
            wb.save(filepath)
            return True
        except Exception as e:
            print(f"Error exporting presensi: {e}")
            return False

    def _export_wilayah_sheet(self, wb: Workbook):
        """Add wilayah sheet to workbook"""
        ws = wb.create_sheet("Wilayah")
        headers = ["ID", "Kode", "Nama", "Tingkat", "Parent_ID", "Is_Aktif", "Sync_ID"]
        ws.append(headers)
        self._style_header(ws, 1, len(headers))

        for idx, w in enumerate(self.session.query(Wilayah).all(), 1):
            self._add_data_row(ws, idx + 1, [
                w.id, w.kode, w.nama, w.tingkat, w.parent_id,
                1 if w.is_aktif else 0, w.sync_id
            ])
        self._auto_width(ws)

    def _export_jenjang_sheet(self, wb: Workbook):
        """Add jenjang sheet to workbook"""
        ws = wb.create_sheet("Jenjang")
        headers = ["ID", "Kode", "Nama", "Usia_Mulai", "Usia_Sampai", "Urutan", "Is_Aktif", "Sync_ID"]
        ws.append(headers)
        self._style_header(ws, 1, len(headers))

        for idx, j in enumerate(self.session.query(Jenjang).all(), 1):
            self._add_data_row(ws, idx + 1, [
                j.id, j.kode, j.nama, j.usia_mulai, j.usia_sampai,
                j.urutan, 1 if j.is_aktif else 0, j.sync_id
            ])
        self._auto_width(ws)

    def _export_jamaah_sheet(self, wb: Workbook):
        """Add jamaah sheet to workbook"""
        ws = wb.create_sheet("Jamaah")
        headers = [
            "ID", "Nama", "Nama_Panggilan", "Jenis_Kelamin", "Tempat_Lahir",
            "Tanggal_Lahir", "Phone", "Email", "Alamat_Lengkap", "Status_Aktif",
            "Status_Pernikahan", "Pasangan_ID", "Tanggal_Menikah", "Ayah_ID", "Ibu_ID",
            "Golongan_Darah", "Pendidikan_Terakhir", "Pekerjaan", "Sync_ID"
        ]
        ws.append(headers)
        self._style_header(ws, 1, len(headers))

        for idx, j in enumerate(self.session.query(Jamaah).all(), 1):
            self._add_data_row(ws, idx + 1, [
                j.id, j.nama, j.nama_panggilan, j.jenis_kelamin, j.tempat_lahir,
                j.tanggal_lahir.isoformat() if j.tanggal_lahir else None,
                j.phone, j.email, j.alamat_lengkap, 1 if j.status_aktif else 0,
                j.status_pernikahan, j.pasangan_id,
                j.tanggal_menikah.isoformat() if j.tanggal_menikah else None,
                j.ayah_id, j.ibu_id, j.golongan_darah, j.pendidikan_terakhir,
                j.pekerjaan, j.sync_id
            ])
        self._auto_width(ws)

    def _export_enrollment_sheet(self, wb: Workbook):
        """Add enrollment sheet to workbook"""
        ws = wb.create_sheet("Enrollment")
        headers = [
            "ID", "Jamaah_ID", "Wilayah_ID", "Jenjang_ID", "Tahun_Ajaran_ID",
            "Status", "Tanggal_Mulai", "Tanggal_Selesai", "Keterangan", "Sync_ID"
        ]
        ws.append(headers)
        self._style_header(ws, 1, len(headers))

        for idx, e in enumerate(self.session.query(Enrollment).all(), 1):
            self._add_data_row(ws, idx + 1, [
                e.id, e.jamaah_id, e.wilayah_id, e.jenjang_id, e.tahun_ajaran_id,
                e.status,
                e.tanggal_mulai.isoformat() if e.tanggal_mulai else None,
                e.tanggal_selesai.isoformat() if e.tanggal_selesai else None,
                e.keterangan, e.sync_id
            ])
        self._auto_width(ws)

    def _export_kurikulum_sheet(self, wb: Workbook):
        """Add kurikulum sheets to workbook"""
        # Bidang Materi
        ws1 = wb.create_sheet("BidangMateri")
        headers1 = ["ID", "Nama", "Urutan", "Is_Aktif", "Sync_ID"]
        ws1.append(headers1)
        self._style_header(ws1, 1, len(headers1))
        for idx, b in enumerate(self.session.query(BidangMateri).all(), 1):
            self._add_data_row(ws1, idx + 1, [
                b.id, b.nama, b.urutan, 1 if b.is_aktif else 0, b.sync_id
            ])
        self._auto_width(ws1)

        # Kategori Materi
        ws2 = wb.create_sheet("KategoriMateri")
        headers2 = ["ID", "Bidang_ID", "Nama", "Kode", "Urutan", "Is_Aktif", "Sync_ID"]
        ws2.append(headers2)
        self._style_header(ws2, 1, len(headers2))
        for idx, k in enumerate(self.session.query(KategoriMateri).all(), 1):
            self._add_data_row(ws2, idx + 1, [
                k.id, k.bidang_id, k.nama, k.kode, k.urutan,
                1 if k.is_aktif else 0, k.sync_id
            ])
        self._auto_width(ws2)

        # Materi Item
        ws3 = wb.create_sheet("MateriItem")
        headers3 = ["ID", "Kategori_ID", "Nama", "Nomor", "Tipe", "Is_Aktif", "Sync_ID"]
        ws3.append(headers3)
        self._style_header(ws3, 1, len(headers3))
        for idx, m in enumerate(self.session.query(MateriItem).all(), 1):
            self._add_data_row(ws3, idx + 1, [
                m.id, m.kategori_id, m.nama, m.nomor, m.tipe,
                1 if m.is_aktif else 0, m.sync_id
            ])
        self._auto_width(ws3)

    def _export_pengajian_sheet(self, wb: Workbook):
        """Add pengajian sheets to workbook"""
        # Jadwal Rutin
        ws1 = wb.create_sheet("JadwalRutin")
        headers1 = ["ID", "Wilayah_ID", "Jenjang_ID", "Nama", "Hari", "Jam", "Is_Aktif", "Sync_ID"]
        ws1.append(headers1)
        self._style_header(ws1, 1, len(headers1))
        for idx, j in enumerate(self.session.query(JadwalRutin).all(), 1):
            self._add_data_row(ws1, idx + 1, [
                j.id, j.wilayah_id, j.jenjang_id, j.nama, j.hari,
                j.jam.isoformat() if j.jam else None, 1 if j.is_aktif else 0, j.sync_id
            ])
        self._auto_width(ws1)

        # Pengajian
        ws2 = wb.create_sheet("Pengajian")
        headers2 = [
            "ID", "Wilayah_ID", "Jenjang_ID", "Jadwal_Rutin_ID", "Tanggal",
            "Waktu_Mulai", "Waktu_Selesai", "Materi_Pokok", "Penyaji_ID",
            "Catatan", "Status", "Sync_ID"
        ]
        ws2.append(headers2)
        self._style_header(ws2, 1, len(headers2))
        for idx, p in enumerate(self.session.query(Pengajian).all(), 1):
            self._add_data_row(ws2, idx + 1, [
                p.id, p.wilayah_id, p.jenjang_id, p.jadwal_rutin_id,
                p.tanggal.isoformat() if p.tanggal else None,
                p.waktu_mulai.isoformat() if p.waktu_mulai else None,
                p.waktu_selesai.isoformat() if p.waktu_selesai else None,
                p.materi_pokok, p.penyaji_id, p.catatan, p.status, p.sync_id
            ])
        self._auto_width(ws2)

    def _export_presensi_sheet(self, wb: Workbook):
        """Add presensi sheet to workbook"""
        ws = wb.create_sheet("Presensi")
        headers = ["ID", "Pengajian_ID", "Jamaah_ID", "Status", "Keterangan", "Sync_ID"]
        ws.append(headers)
        self._style_header(ws, 1, len(headers))
        for idx, p in enumerate(self.session.query(KeaktifanPengajian).all(), 1):
            self._add_data_row(ws, idx + 1, [
                p.id, p.pengajian_id, p.jamaah_id, p.status, p.keterangan, p.sync_id
            ])
        self._auto_width(ws)

    def _export_organisasi_sheet(self, wb: Workbook):
        """Add organisasi sheet to workbook"""
        ws = wb.create_sheet("LimaUnsur")
        headers = ["ID", "Wilayah_ID", "Jabatan", "Jamaah_ID", "Periode_Mulai", "Periode_Selesai", "Is_Aktif", "Sync_ID"]
        ws.append(headers)
        self._style_header(ws, 1, len(headers))
        for idx, l in enumerate(self.session.query(LimaUnsur).all(), 1):
            self._add_data_row(ws, idx + 1, [
                l.id, l.wilayah_id, l.jabatan, l.jamaah_id,
                l.periode_mulai.isoformat() if l.periode_mulai else None,
                l.periode_selesai.isoformat() if l.periode_selesai else None,
                1 if l.is_aktif else 0, l.sync_id
            ])
        self._auto_width(ws)

    # ==================== IMPORT METHODS ====================

    def import_all(self, filepath: str) -> Dict[str, Any]:
        """Import semua data dari file Excel backup"""
        result = {
            'success': False,
            'imported': {},
            'errors': []
        }

        try:
            wb = load_workbook(filepath)

            # Import order matters due to foreign keys
            import_order = [
                ('Wilayah', self._import_wilayah_sheet),
                ('Jenjang', self._import_jenjang_sheet),
                ('Jamaah', self._import_jamaah_sheet),
                ('Enrollment', self._import_enrollment_sheet),
                ('BidangMateri', self._import_bidang_sheet),
                ('KategoriMateri', self._import_kategori_sheet),
                ('MateriItem', self._import_materi_sheet),
                ('JadwalRutin', self._import_jadwal_sheet),
                ('Pengajian', self._import_pengajian_sheet),
                ('Presensi', self._import_presensi_sheet),
                ('LimaUnsur', self._import_limaunsur_sheet),
            ]

            for sheet_name, import_func in import_order:
                if sheet_name in wb.sheetnames:
                    try:
                        count = import_func(wb[sheet_name])
                        result['imported'][sheet_name] = count
                    except Exception as e:
                        result['errors'].append(f"{sheet_name}: {str(e)}")

            self.session.commit()
            result['success'] = len(result['errors']) == 0

        except Exception as e:
            self.session.rollback()
            result['errors'].append(f"General error: {str(e)}")

        return result

    def import_jamaah(self, filepath: str) -> Dict[str, Any]:
        """Import data jamaah dari Excel"""
        result = {'success': False, 'imported': 0, 'errors': [], 'skipped': 0}

        try:
            wb = load_workbook(filepath)
            ws = wb.active

            headers = [cell.value for cell in ws[1]]

            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
                try:
                    if not row[1]:  # Skip empty nama
                        result['skipped'] += 1
                        continue

                    data = dict(zip(headers, row))

                    # Check if exists by sync_id or nama
                    existing = None
                    if data.get('Sync_ID'):
                        existing = self.session.query(Jamaah).filter(
                            Jamaah.sync_id == data['Sync_ID']
                        ).first()

                    if not existing and data.get('Nama'):
                        existing = self.session.query(Jamaah).filter(
                            Jamaah.nama == data['Nama'],
                            Jamaah.tanggal_lahir == self._parse_date(data.get('Tanggal_Lahir'))
                        ).first()

                    if existing:
                        # Update existing
                        self._update_jamaah_from_data(existing, data)
                    else:
                        # Create new
                        jamaah = self._create_jamaah_from_data(data)
                        self.session.add(jamaah)

                    result['imported'] += 1

                except Exception as e:
                    result['errors'].append(f"Row {row_idx}: {str(e)}")

            self.session.commit()
            result['success'] = True

        except Exception as e:
            self.session.rollback()
            result['errors'].append(f"File error: {str(e)}")

        return result

    def import_wilayah(self, filepath: str) -> Dict[str, Any]:
        """Import data wilayah dari Excel"""
        result = {'success': False, 'imported': 0, 'errors': []}

        try:
            wb = load_workbook(filepath)
            ws = wb.active

            headers = [cell.value for cell in ws[1]]

            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
                try:
                    if not row[2]:  # Skip empty nama
                        continue

                    data = dict(zip(headers, row))

                    existing = None
                    if data.get('Sync_ID'):
                        existing = self.session.query(Wilayah).filter(
                            Wilayah.sync_id == data['Sync_ID']
                        ).first()

                    if not existing and data.get('Kode'):
                        existing = self.session.query(Wilayah).filter(
                            Wilayah.kode == data['Kode']
                        ).first()

                    if existing:
                        existing.nama = data.get('Nama', existing.nama)
                        existing.tingkat = data.get('Tingkat', existing.tingkat)
                        existing.is_aktif = bool(data.get('Is_Aktif', 1))
                    else:
                        wilayah = Wilayah(
                            kode=data.get('Kode'),
                            nama=data['Nama'],
                            tingkat=data.get('Tingkat', 'kelompok'),
                            parent_id=data.get('Parent_ID'),
                            is_aktif=bool(data.get('Is_Aktif', 1)),
                            sync_id=data.get('Sync_ID')
                        )
                        self.session.add(wilayah)

                    result['imported'] += 1

                except Exception as e:
                    result['errors'].append(f"Row {row_idx}: {str(e)}")

            self.session.commit()
            result['success'] = True

        except Exception as e:
            self.session.rollback()
            result['errors'].append(f"File error: {str(e)}")

        return result

    def import_kurikulum(self, filepath: str) -> Dict[str, Any]:
        """Import data kurikulum dari Excel"""
        result = {'success': False, 'imported': {}, 'errors': []}

        try:
            wb = load_workbook(filepath)

            # Import Bidang
            if 'BidangMateri' in wb.sheetnames or 'Bidang Materi' in wb.sheetnames:
                sheet_name = 'BidangMateri' if 'BidangMateri' in wb.sheetnames else 'Bidang Materi'
                count = self._import_bidang_sheet(wb[sheet_name])
                result['imported']['BidangMateri'] = count

            # Import Kategori
            if 'KategoriMateri' in wb.sheetnames or 'Kategori Materi' in wb.sheetnames:
                sheet_name = 'KategoriMateri' if 'KategoriMateri' in wb.sheetnames else 'Kategori Materi'
                count = self._import_kategori_sheet(wb[sheet_name])
                result['imported']['KategoriMateri'] = count

            # Import Materi
            if 'MateriItem' in wb.sheetnames or 'Materi Item' in wb.sheetnames:
                sheet_name = 'MateriItem' if 'MateriItem' in wb.sheetnames else 'Materi Item'
                count = self._import_materi_sheet(wb[sheet_name])
                result['imported']['MateriItem'] = count

            self.session.commit()
            result['success'] = True

        except Exception as e:
            self.session.rollback()
            result['errors'].append(str(e))

        return result

    # Helper methods for import
    def _parse_date(self, value) -> Optional[date]:
        """Parse date from various formats"""
        if not value:
            return None
        if isinstance(value, date):
            return value
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, str):
            try:
                return date.fromisoformat(value)
            except:
                try:
                    return datetime.strptime(value, "%Y-%m-%d").date()
                except:
                    return None
        return None

    def _parse_time(self, value) -> Optional[time]:
        """Parse time from various formats"""
        if not value:
            return None
        if isinstance(value, time):
            return value
        if isinstance(value, datetime):
            return value.time()
        if isinstance(value, str):
            try:
                return time.fromisoformat(value)
            except:
                try:
                    return datetime.strptime(value, "%H:%M:%S").time()
                except:
                    try:
                        return datetime.strptime(value, "%H:%M").time()
                    except:
                        return None
        return None

    def _create_jamaah_from_data(self, data: Dict) -> Jamaah:
        """Create Jamaah instance from import data"""
        return Jamaah(
            nama=data.get('Nama') or data.get('nama'),
            nama_panggilan=data.get('Nama_Panggilan') or data.get('nama_panggilan'),
            jenis_kelamin=self._parse_gender(data.get('Jenis_Kelamin') or data.get('jenis_kelamin')),
            tempat_lahir=data.get('Tempat_Lahir') or data.get('tempat_lahir'),
            tanggal_lahir=self._parse_date(data.get('Tanggal_Lahir') or data.get('tanggal_lahir')),
            phone=data.get('Phone') or data.get('phone'),
            email=data.get('Email') or data.get('email'),
            alamat_lengkap=data.get('Alamat_Lengkap') or data.get('alamat') or data.get('Alamat'),
            status_aktif=bool(data.get('Status_Aktif', 1) if data.get('Status_Aktif') != 'Tidak Aktif' else 0),
            status_pernikahan=data.get('Status_Pernikahan') or data.get('status_pernikahan') or 'belum_menikah',
            golongan_darah=data.get('Golongan_Darah') or data.get('golongan_darah'),
            pendidikan_terakhir=data.get('Pendidikan_Terakhir') or data.get('pendidikan') or data.get('Pendidikan'),
            pekerjaan=data.get('Pekerjaan') or data.get('pekerjaan'),
            sync_id=data.get('Sync_ID') or data.get('sync_id')
        )

    def _update_jamaah_from_data(self, jamaah: Jamaah, data: Dict):
        """Update existing Jamaah from import data"""
        if data.get('Nama') or data.get('nama'):
            jamaah.nama = data.get('Nama') or data.get('nama')
        if data.get('Nama_Panggilan') or data.get('nama_panggilan'):
            jamaah.nama_panggilan = data.get('Nama_Panggilan') or data.get('nama_panggilan')
        if data.get('Phone') or data.get('phone'):
            jamaah.phone = data.get('Phone') or data.get('phone')
        if data.get('Email') or data.get('email'):
            jamaah.email = data.get('Email') or data.get('email')
        if data.get('Alamat_Lengkap') or data.get('alamat'):
            jamaah.alamat_lengkap = data.get('Alamat_Lengkap') or data.get('alamat')

    def _parse_gender(self, value) -> Optional[str]:
        """Parse gender from various formats"""
        if not value:
            return None
        value = str(value).lower().strip()
        if value in ['l', 'laki-laki', 'laki', 'male', 'm', 'pria']:
            return 'L'
        if value in ['p', 'perempuan', 'female', 'f', 'wanita']:
            return 'P'
        return value[0].upper() if value else None

    def _import_wilayah_sheet(self, ws) -> int:
        """Import wilayah from sheet"""
        headers = [cell.value for cell in ws[1]]
        count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            data = dict(zip(headers, row))
            if not data.get('Nama'):
                continue
            existing = self.session.query(Wilayah).filter(
                Wilayah.sync_id == data.get('Sync_ID')
            ).first() if data.get('Sync_ID') else None

            if not existing:
                wilayah = Wilayah(
                    kode=data.get('Kode'),
                    nama=data['Nama'],
                    tingkat=data.get('Tingkat', 'kelompok'),
                    parent_id=data.get('Parent_ID'),
                    is_aktif=bool(data.get('Is_Aktif', 1)),
                    sync_id=data.get('Sync_ID')
                )
                self.session.add(wilayah)
                count += 1
        return count

    def _import_jenjang_sheet(self, ws) -> int:
        """Import jenjang from sheet"""
        headers = [cell.value for cell in ws[1]]
        count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            data = dict(zip(headers, row))
            if not data.get('Nama'):
                continue
            existing = self.session.query(Jenjang).filter(
                Jenjang.kode == data.get('Kode')
            ).first()

            if not existing:
                jenjang = Jenjang(
                    kode=data.get('Kode'),
                    nama=data['Nama'],
                    usia_mulai=data.get('Usia_Mulai', 0),
                    usia_sampai=data.get('Usia_Sampai', 100),
                    urutan=data.get('Urutan', 0),
                    is_aktif=bool(data.get('Is_Aktif', 1)),
                    sync_id=data.get('Sync_ID')
                )
                self.session.add(jenjang)
                count += 1
        return count

    def _import_jamaah_sheet(self, ws) -> int:
        """Import jamaah from sheet"""
        headers = [cell.value for cell in ws[1]]
        count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            data = dict(zip(headers, row))
            if not data.get('Nama'):
                continue
            existing = self.session.query(Jamaah).filter(
                Jamaah.sync_id == data.get('Sync_ID')
            ).first() if data.get('Sync_ID') else None

            if not existing:
                jamaah = self._create_jamaah_from_data(data)
                self.session.add(jamaah)
                count += 1
        return count

    def _import_enrollment_sheet(self, ws) -> int:
        """Import enrollment from sheet"""
        headers = [cell.value for cell in ws[1]]
        count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            data = dict(zip(headers, row))
            if not data.get('Jamaah_ID') or not data.get('Wilayah_ID'):
                continue
            existing = self.session.query(Enrollment).filter(
                Enrollment.sync_id == data.get('Sync_ID')
            ).first() if data.get('Sync_ID') else None

            if not existing:
                enrollment = Enrollment(
                    jamaah_id=data['Jamaah_ID'],
                    wilayah_id=data['Wilayah_ID'],
                    jenjang_id=data.get('Jenjang_ID'),
                    tahun_ajaran_id=data.get('Tahun_Ajaran_ID'),
                    status=data.get('Status', 'aktif'),
                    tanggal_mulai=self._parse_date(data.get('Tanggal_Mulai')),
                    tanggal_selesai=self._parse_date(data.get('Tanggal_Selesai')),
                    keterangan=data.get('Keterangan'),
                    sync_id=data.get('Sync_ID')
                )
                self.session.add(enrollment)
                count += 1
        return count

    def _import_bidang_sheet(self, ws) -> int:
        """Import bidang materi from sheet"""
        headers = [cell.value for cell in ws[1]]
        count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            data = dict(zip(headers, row))
            nama = data.get('Nama') or data.get('Nama Bidang')
            if not nama:
                continue
            existing = self.session.query(BidangMateri).filter(
                BidangMateri.nama == nama
            ).first()

            if not existing:
                bidang = BidangMateri(
                    nama=nama,
                    urutan=data.get('Urutan', 0),
                    is_aktif=bool(data.get('Is_Aktif', 1) if data.get('Is_Aktif') != 'Tidak Aktif' else 0),
                    sync_id=data.get('Sync_ID')
                )
                self.session.add(bidang)
                count += 1
        return count

    def _import_kategori_sheet(self, ws) -> int:
        """Import kategori materi from sheet"""
        headers = [cell.value for cell in ws[1]]
        count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            data = dict(zip(headers, row))
            nama = data.get('Nama') or data.get('Nama Kategori')
            if not nama:
                continue
            existing = self.session.query(KategoriMateri).filter(
                KategoriMateri.nama == nama,
                KategoriMateri.bidang_id == data.get('Bidang_ID')
            ).first()

            if not existing:
                kategori = KategoriMateri(
                    bidang_id=data.get('Bidang_ID'),
                    nama=nama,
                    kode=data.get('Kode'),
                    urutan=data.get('Urutan', 0),
                    is_aktif=bool(data.get('Is_Aktif', 1) if data.get('Is_Aktif') != 'Tidak Aktif' else 0),
                    sync_id=data.get('Sync_ID')
                )
                self.session.add(kategori)
                count += 1
        return count

    def _import_materi_sheet(self, ws) -> int:
        """Import materi item from sheet"""
        headers = [cell.value for cell in ws[1]]
        count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            data = dict(zip(headers, row))
            nama = data.get('Nama') or data.get('Nama Materi')
            if not nama:
                continue
            existing = self.session.query(MateriItem).filter(
                MateriItem.nama == nama,
                MateriItem.kategori_id == data.get('Kategori_ID')
            ).first()

            if not existing:
                materi = MateriItem(
                    kategori_id=data.get('Kategori_ID'),
                    nama=nama,
                    nomor=data.get('Nomor'),
                    tipe=data.get('Tipe', 'hafalan'),
                    is_aktif=bool(data.get('Is_Aktif', 1) if data.get('Is_Aktif') != 'Tidak Aktif' else 0),
                    sync_id=data.get('Sync_ID')
                )
                self.session.add(materi)
                count += 1
        return count

    def _import_jadwal_sheet(self, ws) -> int:
        """Import jadwal rutin from sheet"""
        headers = [cell.value for cell in ws[1]]
        count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            data = dict(zip(headers, row))
            if not data.get('Hari'):
                continue
            existing = self.session.query(JadwalRutin).filter(
                JadwalRutin.sync_id == data.get('Sync_ID')
            ).first() if data.get('Sync_ID') else None

            if not existing:
                jadwal = JadwalRutin(
                    wilayah_id=data.get('Wilayah_ID'),
                    jenjang_id=data.get('Jenjang_ID'),
                    nama=data.get('Nama'),
                    hari=data['Hari'],
                    jam=self._parse_time(data.get('Jam')),
                    is_aktif=bool(data.get('Is_Aktif', 1)),
                    sync_id=data.get('Sync_ID')
                )
                self.session.add(jadwal)
                count += 1
        return count

    def _import_pengajian_sheet(self, ws) -> int:
        """Import pengajian from sheet"""
        headers = [cell.value for cell in ws[1]]
        count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            data = dict(zip(headers, row))
            if not data.get('Tanggal'):
                continue
            existing = self.session.query(Pengajian).filter(
                Pengajian.sync_id == data.get('Sync_ID')
            ).first() if data.get('Sync_ID') else None

            if not existing:
                pengajian = Pengajian(
                    wilayah_id=data.get('Wilayah_ID'),
                    jenjang_id=data.get('Jenjang_ID'),
                    jadwal_rutin_id=data.get('Jadwal_Rutin_ID'),
                    tanggal=self._parse_date(data['Tanggal']),
                    waktu_mulai=self._parse_time(data.get('Waktu_Mulai')),
                    waktu_selesai=self._parse_time(data.get('Waktu_Selesai')),
                    materi_pokok=data.get('Materi_Pokok'),
                    penyaji_id=data.get('Penyaji_ID'),
                    catatan=data.get('Catatan'),
                    status=data.get('Status', 'terlaksana'),
                    sync_id=data.get('Sync_ID')
                )
                self.session.add(pengajian)
                count += 1
        return count

    def _import_presensi_sheet(self, ws) -> int:
        """Import presensi from sheet"""
        headers = [cell.value for cell in ws[1]]
        count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            data = dict(zip(headers, row))
            if not data.get('Pengajian_ID') or not data.get('Jamaah_ID'):
                continue
            existing = self.session.query(KeaktifanPengajian).filter(
                KeaktifanPengajian.sync_id == data.get('Sync_ID')
            ).first() if data.get('Sync_ID') else None

            if not existing:
                presensi = KeaktifanPengajian(
                    pengajian_id=data['Pengajian_ID'],
                    jamaah_id=data['Jamaah_ID'],
                    status=data.get('Status', 'hadir'),
                    keterangan=data.get('Keterangan'),
                    sync_id=data.get('Sync_ID')
                )
                self.session.add(presensi)
                count += 1
        return count

    def _import_limaunsur_sheet(self, ws) -> int:
        """Import lima unsur from sheet"""
        headers = [cell.value for cell in ws[1]]
        count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            data = dict(zip(headers, row))
            if not data.get('Wilayah_ID') or not data.get('Jabatan'):
                continue
            existing = self.session.query(LimaUnsur).filter(
                LimaUnsur.sync_id == data.get('Sync_ID')
            ).first() if data.get('Sync_ID') else None

            if not existing:
                unsur = LimaUnsur(
                    wilayah_id=data['Wilayah_ID'],
                    jabatan=data['Jabatan'],
                    jamaah_id=data.get('Jamaah_ID'),
                    periode_mulai=self._parse_date(data.get('Periode_Mulai')),
                    periode_selesai=self._parse_date(data.get('Periode_Selesai')),
                    is_aktif=bool(data.get('Is_Aktif', 1)),
                    sync_id=data.get('Sync_ID')
                )
                self.session.add(unsur)
                count += 1
        return count

    def get_template(self, module: str, filepath: str) -> bool:
        """Generate template Excel untuk import"""
        try:
            wb = Workbook()
            ws = wb.active

            templates = {
                'jamaah': [
                    "Nama", "Nama_Panggilan", "Jenis_Kelamin", "Tempat_Lahir",
                    "Tanggal_Lahir", "Phone", "Email", "Alamat_Lengkap",
                    "Status_Pernikahan", "Golongan_Darah", "Pendidikan_Terakhir", "Pekerjaan"
                ],
                'wilayah': ["Kode", "Nama", "Tingkat", "Parent_ID"],
                'kurikulum': ["Bidang", "Kategori", "Nomor", "Nama_Materi", "Tipe"],
            }

            if module in templates:
                ws.title = f"Template {module.title()}"
                ws.append(templates[module])
                self._style_header(ws, 1, len(templates[module]))

                # Add sample row
                if module == 'jamaah':
                    ws.append(["Ahmad", "Amad", "L", "Sorong", "2010-01-15", "081234567890", "", "Jl. Contoh No. 1", "belum_menikah", "O", "SD", "Pelajar"])
                elif module == 'wilayah':
                    ws.append(["SRG-001", "Kelompok Contoh", "kelompok", ""])
                elif module == 'kurikulum':
                    ws.append(["Al-Quran", "Juz 30", "1", "Al-Fatihah", "hafalan"])

                self._auto_width(ws)

            wb.save(filepath)
            return True
        except Exception as e:
            print(f"Error creating template: {e}")
            return False
